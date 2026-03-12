from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any

from sqlalchemy import inspect
from sqlalchemy.sql import text

from ..extensions import db


def export_db_to_json() -> dict[str, Any]:
    payload: dict[str, Any] = {"exported_at": datetime.utcnow().isoformat() + "Z", "tables": {}}
    insp = inspect(db.engine)
    for table_name in insp.get_table_names():
        rows = db.session.execute(text(f"SELECT * FROM `{table_name}`")).mappings().all()
        payload["tables"][table_name] = [dict(r) for r in rows]
    return payload


def export_db_to_xlsx() -> io.BytesIO:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except Exception as e:  # pragma: no cover
        raise RuntimeError("Excel 导出依赖未就绪：请先安装 openpyxl。") from e

    insp = inspect(db.engine)
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    def cell_value(v: Any):
        if v is None:
            return None
        if isinstance(v, (dict, list)):
            return json.dumps(v, ensure_ascii=False, default=str)
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v

    def _read_rows(table_name: str):
        return db.session.execute(text(f"SELECT * FROM `{table_name}`")).mappings().all()

    def _auto_format(ws, columns: list[str]):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(columns, start=1):
            max_len = len(str(col))
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx, max_row=min(ws.max_row, 2000)):
                v = row[0].value
                if v is None:
                    continue
                max_len = max(max_len, min(len(str(v)), 60))
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max_len + 2, 70)

    def _write_sheet(title: str, columns: list[str], rows: list[dict[str, Any]]):
        ws = wb.create_sheet(title=title[:31])
        ws.append(columns)
        for col_idx in range(1, len(columns) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = header_font
            c.alignment = header_alignment

        for row in rows:
            ws.append([cell_value(row.get(c)) for c in columns])

        _auto_format(ws, columns)

    exported_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    table_names = sorted(insp.get_table_names())

    # 00: export guide
    ws0 = wb.create_sheet(title="00_导出说明")
    ws0.append(["导出时间", exported_at])
    ws0.append(["导出类型", "Excel（用于查阅/审计，不用于数据恢复）"])
    ws0.append(["恢复建议", "请使用 backup.json 执行恢复；Excel 仅作展示/归档"])
    ws0.append(["说明", "本文件包含：业务概览、业务明细、原始表快照（RAW_前缀）"])
    ws0.append([])
    ws0.append(["工作表", "用途"])
    ws0.append(["01_业务概览", "关键业务统计摘要"])
    ws0.append(["02_用户清单", "用户账户状态（密码已脱敏）"])
    ws0.append(["03_化学品清单", "危化品基础参数"])
    ws0.append(["04_储存记录", "企业储存信息明细"])
    ws0.append(["05_评估结果", "辨识评估结果与审核状态"])
    ws0.append(["RAW_*", "系统原始表快照，便于技术排查"])
    for col_idx in range(1, 3):
        c = ws0.cell(row=6, column=col_idx)
        c.font = header_font
        c.alignment = header_alignment
    ws0.freeze_panes = "A7"
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 76

    # 01: business overview
    users = _read_rows("users") if "users" in table_names else []
    chemicals = _read_rows("chemicals") if "chemicals" in table_names else []
    storages = _read_rows("storage_records") if "storage_records" in table_names else []
    results = _read_rows("evaluation_results") if "evaluation_results" in table_names else []

    ws1 = wb.create_sheet(title="01_业务概览")
    ws1.append(["指标", "数值", "说明"])
    metrics = [
        ("用户总数", len(users), "系统账号数量"),
        ("管理员数量", sum(1 for r in users if r.get("role") == "admin"), "管理员角色账号"),
        ("启用账号", sum(1 for r in users if r.get("is_active") in (True, 1)), "可登录账号"),
        ("化学品数量", len(chemicals), "危化品参数库"),
        ("储存记录数量", len(storages), "企业录入储存明细"),
        ("评估结果数量", len(results), "已生成辨识评估结果"),
        ("待审核结果", sum(1 for r in results if r.get("status") == "pending"), "待管理员处理"),
        ("已通过结果", sum(1 for r in results if r.get("status") == "approved"), "审核通过"),
        ("已驳回结果", sum(1 for r in results if r.get("status") == "rejected"), "审核驳回"),
        ("重大危险源结果", sum(1 for r in results if r.get("is_major_hazard") in (True, 1)), "判定为重大危险源"),
    ]
    for m in metrics:
        ws1.append(list(m))
    for col_idx in range(1, 4):
        c = ws1.cell(row=1, column=col_idx)
        c.font = header_font
        c.alignment = header_alignment
    ws1.freeze_panes = "A2"
    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 12
    ws1.column_dimensions["C"].width = 42

    # 02-05: business-friendly sheets
    if users:
        rows = []
        for r in users:
            rows.append(
                {
                    "用户名": r.get("username"),
                    "角色": "管理员" if r.get("role") == "admin" else "普通用户",
                    "状态": "启用" if r.get("is_active") in (True, 1) else "停用",
                    "企业名称": r.get("company_name"),
                    "联系人": r.get("contact_name"),
                    "联系电话": r.get("phone"),
                    "地址": r.get("address"),
                    "创建时间": cell_value(r.get("created_at")),
                }
            )
        _write_sheet("02_用户清单", list(rows[0].keys()), rows)

    if chemicals:
        rows = []
        for r in chemicals:
            rows.append(
                {
                    "名称": r.get("name"),
                    "类别": r.get("category"),
                    "CAS号": r.get("cas_no"),
                    "临界量": r.get("critical_quantity"),
                    "单位": r.get("unit"),
                    "类别符号": r.get("hazard_category_symbol"),
                    "β": r.get("beta"),
                    "标准来源": r.get("source_standard"),
                }
            )
        _write_sheet("03_化学品清单", list(rows[0].keys()), rows)

    if storages:
        rows = []
        for r in storages:
            rows.append(
                {
                    "企业": r.get("enterprise_name"),
                    "化学品ID": r.get("chemical_id"),
                    "储存量": r.get("amount"),
                    "单位": r.get("unit"),
                    "单元名称": r.get("unit_name"),
                    "单元类型": r.get("unit_type"),
                    "场所": r.get("location"),
                    "储存方式": r.get("storage_method"),
                    "备注": r.get("properties"),
                    "更新时间": cell_value(r.get("updated_at")),
                }
            )
        _write_sheet("04_储存记录", list(rows[0].keys()), rows)

    if results:
        rows = []
        for r in results:
            rows.append(
                {
                    "企业": r.get("enterprise_name"),
                    "R值": r.get("r_value"),
                    "是否重大危险源": "是" if r.get("is_major_hazard") in (True, 1) else "否",
                    "GB等级": r.get("level_by_gb") or r.get("level"),
                    "审核状态": r.get("status"),
                    "审核备注": r.get("review_remark"),
                    "评估时间": cell_value(r.get("created_at")),
                }
            )
        _write_sheet("05_评估结果", list(rows[0].keys()), rows)

    # RAW: full table snapshots for troubleshooting.
    for table_name in table_names:
        rows = _read_rows(table_name)
        if rows:
            columns = list(rows[0].keys())
        else:
            columns = [c["name"] for c in insp.get_columns(table_name)]

        ws = wb.create_sheet(title=f"RAW_{table_name}"[:31])
        ws.append(columns)
        for col_idx in range(1, len(columns) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = header_font
            c.alignment = header_alignment

        for r in rows:
            row_dict = dict(r)
            if table_name == "users" and "password_hash" in row_dict:
                row_dict["password_hash"] = "***"
            ws.append([cell_value(row_dict.get(c)) for c in columns])

        _auto_format(ws, columns)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def import_db_from_json(payload: dict[str, Any], *, truncate_first: bool = True) -> None:
    tables: dict[str, list[dict[str, Any]]] = payload.get("tables", {}) or {}
    insp = inspect(db.engine)
    insert_order = _toposort_tables(list(tables.keys()), insp)
    delete_order = list(reversed(insert_order))

    def normalize(v: Any):
        if v is None:
            return None
        if isinstance(v, (dict, list)):
            return json.dumps(v, ensure_ascii=False, default=str)
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v

    # IMPORTANT: FOREIGN_KEY_CHECKS is per-connection in MySQL, so pin one connection.
    with db.engine.begin() as conn:
        _set_fk_checks_conn(conn, False)

        if truncate_first:
            for table_name in delete_order:
                conn.execute(text(f"DELETE FROM `{table_name}`"))

        for table_name in insert_order:
            rows = tables.get(table_name, [])
            if not rows:
                continue
            cols = list(rows[0].keys())
            col_list = ", ".join([f"`{c}`" for c in cols])
            val_list = ", ".join([f":{c}" for c in cols])
            stmt = text(f"INSERT INTO `{table_name}` ({col_list}) VALUES ({val_list})")
            cooked = [{k: normalize(v) for k, v in r.items()} for r in rows]
            conn.execute(stmt, cooked)

        _set_fk_checks_conn(conn, True)


def _set_fk_checks_conn(conn, enabled: bool) -> None:
    try:
        conn.execute(text(f"SET FOREIGN_KEY_CHECKS={'1' if enabled else '0'}"))
    except Exception:
        return


def _toposort_tables(table_names: list[str], insp) -> list[str]:
    # Parents before children (insert order), reverse for delete order.
    table_set = set(table_names)
    deps: dict[str, set[str]] = {t: set() for t in table_names}
    rdeps: dict[str, set[str]] = {t: set() for t in table_names}
    try:
        for t in table_names:
            for fk in insp.get_foreign_keys(t) or []:
                ref = fk.get("referred_table")
                if ref and ref in table_set:
                    deps[t].add(ref)
                    rdeps[ref].add(t)
    except Exception:
        return table_names

    ready = [t for t in table_names if not deps[t]]
    out: list[str] = []
    while ready:
        n = ready.pop(0)
        out.append(n)
        for child in list(rdeps[n]):
            deps[child].discard(n)
            if not deps[child]:
                ready.append(child)
    remaining = [t for t in table_names if t not in out]
    return out + remaining


def dumps_json(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str)


def loads_json(text_in: str) -> dict[str, Any]:
    return json.loads(text_in)
